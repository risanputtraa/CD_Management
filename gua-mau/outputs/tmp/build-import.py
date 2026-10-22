import json
import math
import re
from collections import Counter
from datetime import datetime

import openpyxl


SOURCE = "outputs/tmp/jenius-source.xlsx"
wb = openpyxl.load_workbook(SOURCE, data_only=True, read_only=False)

folders = []
tasks = []
folder_key_counts = Counter()
folder_ids_by_key = {}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip().replace("\uFFFD", "×")
        return value or None
    return value


def date_value(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="milliseconds")
    text = str(value).strip()
    m = re.match(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}:\d{2})(?:\.(\d+))?$", text)
    if m:
        frac = (m.group(3) or "")[:3].ljust(3, "0")
        return f"{m.group(1)}T{m.group(2)}.{frac}"
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).isoformat(timespec="milliseconds")
        except ValueError:
            pass
    return None


def number(value):
    if value in (None, ""):
        return None
    try:
        n = float(value)
        return int(n) if n.is_integer() else None
    except (TypeError, ValueError):
        return None


def status_value(value, started_at):
    s = str(value or "").strip().lower()
    if s in ("revision", "ongoing", "on going"):
        return "Pending"
    if s == "done":
        return "Done"
    if s == "pending":
        return "Pending"
    if s == "open":
        return "Open"
    if s in ("canceled", "cancelled"):
        return "Canceled"
    return "Pending" if started_at else "Open"


def add_folder(parent_id, name):
    name = str(clean(name) or "Unnamed")
    key = (parent_id or "", name.casefold())
    folder_key_counts[key] += 1
    if folder_key_counts[key] > 1:
        name = f"{name} ({folder_key_counts[key]})"
    folder_id = f"folder:{len(folders)+1}"
    folders.append({"folder_id": folder_id, "parent_folder_id": parent_id, "folder_name": name})
    folder_ids_by_key[key] = folder_id
    return folder_id


def add_task(folder_id, row, stage):
    task_type, level_col, pic_col, start_col, status_col, finish_col, duration_col = stage
    level = number(row[level_col]) if level_col is not None and level_col < len(row) else None
    pic = clean(row[pic_col]) if pic_col is not None and pic_col < len(row) else None
    started = date_value(row[start_col]) if start_col is not None and start_col < len(row) else None
    finished = date_value(row[finish_col]) if finish_col is not None and finish_col < len(row) else None
    status = status_value(row[status_col] if status_col is not None and status_col < len(row) else None, started)
    elapsed = 0
    if duration_col is not None and duration_col < len(row):
        try:
            hours = float(row[duration_col] or 0)
            if math.isfinite(hours) and 0 <= hours <= 100000:
                elapsed = round(hours * 3600)
        except (TypeError, ValueError):
            pass
    tasks.append({
        "folder_id": folder_id,
        "task_name": task_type,
        "task_type_name": task_type,
        "difficulty_level": level,
        "pic_name": str(pic) if pic else None,
        "started_at": started,
        "status": status,
        "completed_at": finished,
        "elapsed_seconds": elapsed,
    })


def add_entity(parent_id, name, row, stages, code=None):
    folder_id = add_folder(parent_id, name)
    for stage in stages:
        add_task(folder_id, row, stage)
    return folder_id


# CA scenarios: preserve the source hierarchy system > category > scenario.
ca = wb["VR - CA"]
ca_root = add_folder(None, "CA-Scenario")
system_id = category_id = None
ca_scenarios = 0
ca_stage = [
    ("Materi Deskriptif atau Naratif Gameplay", 8, 9, 10, 11, 12, 15),
    ("QC Materi Deskriptif atau Naratif Gameplay", 29, 30, 31, 32, 33, None),
    ("Visual Brief", 36, 37, 38, 39, 40, 43),
    ("QC Visual Brief", None, 57, 58, 59, 60, None),
    ("Storyboard Gameplay", 63, 64, 65, 66, 67, 70),
    ("QC Storyboard Gameplay", None, 84, 85, 86, 87, None),
    ("QC Model 3D Hardsurface", None, 90, 91, 92, 93, None),
]
for row in ca.iter_rows(min_row=3, values_only=True):
    marker = clean(row[0])
    title = clean(row[1])
    if not title:
        continue
    if isinstance(marker, str) and marker in ("I", "II", "III"):
        system_id = add_folder(ca_root, title)
        category_id = None
    elif isinstance(marker, str) and marker in tuple("ABCDEF") and system_id:
        category_id = add_folder(system_id, title)
    elif isinstance(marker, (int, float)) and category_id:
        add_entity(category_id, title, row, ca_stage, clean(row[2]))
        ca_scenarios += 1


def process_simple_sheet(sheet_name, root_name, key_col, stages):
    sheet = wb[sheet_name]
    root = add_folder(None, root_name)
    count = 0
    for row in sheet.iter_rows(min_row=3, values_only=True):
        name = clean(row[key_col]) if key_col < len(row) else None
        if not name:
            continue
        add_entity(root, str(name), row, stages)
        count += 1
    return count


product_count = process_simple_sheet("VR - 3D - Product", "3D-Product", 2, [
    ("Modelling Product Content", 10, 11, 12, 13, 14, 17),
    ("QC Model 3D Hardsurface", None, 31, 32, 33, 34, None),
    ("Texturing Product Content", 37, 38, 39, 40, 41, 44),
    ("QC Texture 3D Hardsurface", None, 58, 59, 60, 61, None),
    ("Prefabs Making (Product Content/Hardsurface) (Unity)", None, 65, 66, 67, 68, 71),
    ("QC Prefabs Product Content Hardsurface (Unity)", None, 85, 86, 87, 88, None),
])

envi_count = process_simple_sheet("VR - 3D - Envi", "3D-Envi", 1, [
    ("Modelling Environment", 9, 10, 11, 12, 13, 16),
    ("QC Model 3D Environment", None, 30, 31, 32, 33, None),
    ("__TEXTURE_ENVIRONMENT__", 36, 38, 39, 40, 41, 44),
    ("QC Texture 3D Environment", None, 58, 59, 60, 61, None),
    ("Environment 3D Assembly", 64, 65, 66, 67, 68, 71),
    ("QC Assembly 3D Environment", None, 85, 86, 87, 88, None),
])

# Rewrite the texture type based on the Interior/Exterior field in the source.
envi_root = next(x["folder_id"] for x in folders if x["folder_name"] == "3D-Envi")
for row in wb["VR - 3D - Envi"].iter_rows(min_row=3, values_only=True):
    name = clean(row[1])
    if not name:
        continue
    entity_folder = next((x["folder_id"] for x in folders if x["parent_folder_id"] == envi_root and x["folder_name"] == name), None)
    if entity_folder:
        task = next((t for t in tasks if t["folder_id"] == entity_folder and t["task_name"] == "__TEXTURE_ENVIRONMENT__"), None)
        if task:
            env_type = str(clean(row[37] if len(row) > 37 else row[3]) or clean(row[3]) or "Interior").lower()
            task_name = "Texturing Environment Exterior" if "exterior" in env_type else "Texturing Environment Interior"
            task["task_name"] = task["task_type_name"] = task_name

character_count = process_simple_sheet("VR - 3D - Character", "3D-Character", 1, [
    ("Modelling Character on CC", 10, 11, 12, 13, 14, 17),
    ("QC Model 3D Character", None, 31, 32, 33, 34, None),
    ("Texturing Clothe''s Skin", 37, 39, 40, 41, 42, 45),
    ("QC Texture 3D Character", None, 59, 60, 61, 62, None),
    ("Character 3D Cloth Assembly", 65, 66, 67, 68, 69, 72),
    ("QC Assembly 3D Character", None, 86, 87, 88, 89, None),
    ("Character Cut and Weightpaint", None, 92, 93, 94, 95, 98),
    ("QC Assembly 3D Character", None, 112, 113, 114, 115, None),
    ("Prefabs Making (Character) (Unity)", None, 118, 119, 120, 121, 124),
    ("QC Prefabs Character (Unity)", None, 138, 139, 140, 141, None),
])

animation_count = process_simple_sheet("VR - 3D - ANIMATION", "3D-Animation", 1, [
    ("Character Animation", 10, 11, 12, 13, 14, 17),
    ("QC Model 3D Character", None, 31, 32, 33, 34, None),
    ("Mechanical Animation", 37, 38, 39, 40, 41, 44),
    ("QC Animation 3D Mechanical", None, 58, 59, 60, 61, None),
    ("QC Assembly 3D Character", None, 63, 64, 65, 66, None),
    ("Camera Cinematic (Unity)", 7, 69, 70, 71, 72, 75),
    ("QC Animation 3D Camera Cinematic (Unity)", None, 89, 90, 91, 92, None),
])

assy_count = process_simple_sheet("VR - 3D - Assy & Prefabs CF", "3D-Assy & Prefabs CF", 2, [
    ("Mechanical 3D Scenario Assembly", 7, 8, 9, 10, 11, 14),
    ("QC Assembly 3D Mechanical", None, 28, 29, 30, 31, None),
    ("Prefabs Making (Product Content/Hardsurface) (Unity)", 33, 34, 35, 36, 37, 40),
    ("QC Prefabs Scenario (Unity)", None, 54, 55, 56, 57, None),
])

deliverables = {
    "Skenario Gameplay": ca_scenarios,
    "3D Product Content/Envi": product_count,
    "3D Environment": envi_count,
    "3D Character": character_count,
    "3D Animasi": animation_count,
    "3D Game Object Unity": assy_count,
}

out = {
    "project_name": "Jenius VR SMK Madiun",
    "product_type_name": "Jenius VR",
    "start_date": "2026-08-14",
    "deadline": "2027-05-31",
    "status": "On Going",
    "folders": folders,
    "tasks": tasks,
    "deliverables": deliverables,
    "counts": {
        "ca_scenarios": ca_scenarios,
        "products": product_count,
        "environments": envi_count,
        "characters": character_count,
        "animations": animation_count,
        "assy_scenarios": assy_count,
        "folders": len(folders),
        "tasks": len(tasks),
        "statuses": dict(Counter(t["status"] for t in tasks)),
        "source_pic_names": sorted({t["pic_name"] for t in tasks if t["pic_name"]}),
        "task_type_names": sorted({t["task_type_name"] for t in tasks}),
    },
}
with open("outputs/tmp/jenius-import.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=2)
print(json.dumps(out["counts"], ensure_ascii=False, indent=2))
print("DELIVERABLES", json.dumps(deliverables, ensure_ascii=False))
